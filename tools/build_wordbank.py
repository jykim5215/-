#!/usr/bin/env python3
"""
엑셀 단어장(Day / 단어 / 뜻)을 VocaCard 안드로이드 앱의 아카이브 에셋으로 변환한다.

출력: android/app/src/main/assets/wordbank.json

JSON 스키마
{
  "version": 1,
  "source": "voca_30days",
  "days": [
    {
      "id": "day1",
      "index": 1,
      "title": "Day 1",
      "words": [
        {
          "id": "day1-0001",
          "word": "resume",
          "meaning": "이력서",           # 원본 뜻(줄바꿈 유지)
          "senses": ["이력서"]           # 다의어 분리 결과
        }
      ]
    }
  ]
}

사용법:  python3 tools/build_wordbank.py <input.xlsx> [output.json]
"""
from __future__ import annotations

import json
import re
import sys
from collections import OrderedDict
from pathlib import Path

import openpyxl

SENSE_SPLIT = re.compile(r"[\n;；]+")


def split_senses(meaning: str) -> list[str]:
    """다의어 뜻을 개별 sense 로 분리한다. 쉼표는 동의어 나열이므로 자르지 않는다."""
    parts = [p.strip(" ,·") for p in SENSE_SPLIT.split(meaning)]
    return [p for p in parts if p]


def normalize(text: object) -> str:
    return re.sub(r"[ \t ]+", " ", str(text)).strip()


def build(xlsx_path: Path, sheet_name: str = "Sheet1") -> dict:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.worksheets[0]

    days: "OrderedDict[str, list]" = OrderedDict()
    for row in ws.iter_rows(min_row=2, values_only=True):
        day, word, meaning = row[0], row[1], row[2]
        if not word or not meaning:
            continue
        day_id = normalize(day).lower().replace(" ", "") or "day0"
        days.setdefault(day_id, []).append((normalize(word), normalize(meaning)))

    def day_order(day_id: str) -> int:
        m = re.search(r"(\d+)", day_id)
        return int(m.group(1)) if m else 0

    out_days = []
    for day_id in sorted(days, key=day_order):
        idx = day_order(day_id)
        words = []
        for i, (word, meaning) in enumerate(days[day_id], start=1):
            words.append(
                {
                    "id": f"{day_id}-{i:04d}",
                    "word": word,
                    "meaning": meaning,
                    "senses": split_senses(meaning),
                }
            )
        out_days.append(
            {"id": day_id, "index": idx, "title": f"Day {idx}", "words": words}
        )

    return {"version": 1, "source": "voca_30days", "days": out_days}


def main() -> int:
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    src = Path(sys.argv[1])
    dst = Path(
        sys.argv[2]
        if len(sys.argv) > 2
        else "android/app/src/main/assets/wordbank.json"
    )
    data = build(src)
    dst.parent.mkdir(parents=True, exist_ok=True)
    dst.write_text(
        json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8"
    )
    total = sum(len(d["words"]) for d in data["days"])
    poly = sum(1 for d in data["days"] for w in d["words"] if len(w["senses"]) > 1)
    print(f"days={len(data['days'])} words={total} polysemous={poly} -> {dst}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
