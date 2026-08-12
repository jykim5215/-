"""품목 단위(458개) 가중치 적재.

품목 가중치는 KOSIS API로 제공되지 않는다. 국가데이터처 통계설명자료(승인번호
101007)에 첨부된 「지출목적별 품목 및 가중치(2000년~2020년).xlsx」를 1회
내려받아 DB에 적재한다. 갱신 주기가 2~3년이라 수동 갱신으로 충분하다.

파일을 아직 확보하지 못했으면 12대 분류 값으로 시드한다(seed_divisions).
그 경우 계산은 12분류 단위로 이뤄지며, 출처에 그 사실이 남는다.

★ 기준개편 대응: 품목 코드는 2026-12 개편에서 458 -> 455개로 바뀐다(신규 10,
제외 13). 그래서 매핑 테이블에 valid_from/valid_to 를 두고, 가중치 행에는
base_year 를 반드시 붙인다.
"""

from __future__ import annotations

import sqlite3
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable, Sequence

from ..errors import DataUnavailable
from ..provenance import KST
from .official import DIVISIONS, OFFICIAL_WEIGHTS_2022, OFFICIAL_WEIGHT_SOURCE, resolve_division


@dataclass(frozen=True)
class ItemWeight:
    """품목 하나의 가중치."""

    item_code: str
    item_name: str
    weight: float
    division_code: str
    base_year: int


#: 엑셀에서 열을 찾을 때 쓰는 머리글 후보. 파일 판마다 표기가 조금씩 다르다.
_ITEM_HEADERS = ("품목", "품목명", "항목", "지출목적")
_WEIGHT_HEADERS = ("가중치", "가중값", "weight")


def load_from_xlsx(
    path: str | Path,
    *,
    base_year: int,
    sheet_name: str | None = None,
    weight_column: str | None = None,
) -> list[ItemWeight]:
    """통계설명자료 첨부 엑셀에서 품목 가중치를 읽는다.

    파일 판마다 열 구성이 달라서 머리글을 탐색해 열 위치를 잡는다. 자동 탐색이
    실패하면 weight_column 으로 직접 지정한다(예: '2020').

    Raises:
        DataUnavailable: 파일이 없거나 머리글을 찾지 못한 경우.
    """
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - 의존성 누락은 설치로 해결
        raise DataUnavailable("openpyxl 이 필요합니다: pip install openpyxl") from exc

    path = Path(path)
    if not path.exists():
        raise DataUnavailable(
            f"가중치 엑셀을 찾을 수 없습니다: {path}. "
            "국가데이터처 통계설명자료(승인번호 101007)의 "
            "「지출목적별 품목 및 가중치」 파일을 내려받아 지정하세요."
        )

    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
    rows = [list(r) for r in worksheet.iter_rows(values_only=True)]
    workbook.close()

    header_index, item_col, weight_col = _locate_columns(rows, weight_column)

    items: list[ItemWeight] = []
    current_division: str | None = None
    for row in rows[header_index + 1 :]:
        name = _text(row[item_col]) if item_col < len(row) else ""
        if not name:
            continue

        # 12대 분류 행은 그 아래 품목들의 소속을 알려주는 머리 행이다.
        division = resolve_division(name)
        if division is not None:
            current_division = division.code
            continue

        weight = _number(row[weight_col]) if weight_col < len(row) else None
        if weight is None or current_division is None:
            continue

        items.append(
            ItemWeight(
                item_code=_item_code(current_division, len(items)),
                item_name=name,
                weight=weight,
                division_code=current_division,
                base_year=base_year,
            )
        )

    if not items:
        raise DataUnavailable(
            f"{path.name} 에서 품목 가중치를 읽지 못했습니다. "
            "sheet_name/weight_column 을 지정해 보세요."
        )
    return items


#: 머리글로 인정할 셀의 최대 길이. "지출목적별 품목 및 가중치" 같은 제목 행이
#: 품목 열과 가중치 열 양쪽에 걸려 같은 열로 잡히는 것을 막는다.
_MAX_HEADER_LEN = 12


def _locate_columns(rows: Sequence[Sequence], weight_column: str | None) -> tuple[int, int, int]:
    """머리글 행과 품목명/가중치 열 위치를 찾는다.

    품목 열과 가중치 열은 반드시 서로 다른 열이어야 한다.
    """
    for index, row in enumerate(rows[:30]):
        cells = [_text(c) for c in row]
        item_col = _find(cells, _ITEM_HEADERS)
        if item_col is None:
            continue
        if weight_column:
            weight_col = next(
                (i for i, c in enumerate(cells) if c == weight_column and i != item_col),
                None,
            )
        else:
            weight_col = _find(cells, _WEIGHT_HEADERS, exclude=item_col)
        if weight_col is not None and weight_col != item_col:
            return index, item_col, weight_col
    raise DataUnavailable(
        "엑셀에서 '품목'/'가중치' 머리글을 찾지 못했습니다. "
        "sheet_name 또는 weight_column 을 직접 지정하세요."
    )


def _find(
    cells: Sequence[str], candidates: Iterable[str], *, exclude: int | None = None
) -> int | None:
    """머리글 후보를 찾는다. 정확히 일치하는 셀을 우선한다."""
    candidates = tuple(candidates)
    for index, cell in enumerate(cells):
        if index != exclude and cell in candidates:
            return index
    for index, cell in enumerate(cells):
        if index == exclude or len(cell) > _MAX_HEADER_LEN:
            continue
        if any(candidate in cell for candidate in candidates):
            return index
    return None


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _number(value: object) -> float | None:
    if value is None:
        return None
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError:
        return None


def _item_code(division_code: str, ordinal: int) -> str:
    return f"{division_code}-{ordinal:03d}"


# ---------------------------------------------------------------------------
# DB 적재
# ---------------------------------------------------------------------------


def persist(
    conn: sqlite3.Connection,
    items: Sequence[ItemWeight],
    *,
    source: str,
    valid_from: str,
    valid_to: str | None = None,
) -> int:
    """품목 가중치와 매핑을 DB에 넣는다. 매핑에는 유효기간이 붙는다."""
    conn.executemany(
        "INSERT OR REPLACE INTO item_weights "
        "(item_code, item_name, weight, base_year, source) VALUES (?, ?, ?, ?, ?)",
        [(i.item_code, i.item_name, i.weight, i.base_year, source) for i in items],
    )
    conn.executemany(
        "INSERT OR REPLACE INTO item_mapping "
        "(item_code, item_name, division_code, base_year, valid_from, valid_to) "
        "VALUES (?, ?, ?, ?, ?, ?)",
        [
            (i.item_code, i.item_name, i.division_code, i.base_year, valid_from, valid_to)
            for i in items
        ],
    )
    conn.commit()
    return len(items)


def seed_divisions(conn: sqlite3.Connection, *, base_year: int = 2020) -> int:
    """품목 엑셀을 아직 못 구했을 때 12대 분류 값으로 시드한다."""
    items = [
        ItemWeight(
            item_code=division.code,
            item_name=division.name,
            weight=OFFICIAL_WEIGHTS_2022[division.code],
            division_code=division.code,
            base_year=base_year,
        )
        for division in DIVISIONS
    ]
    return persist(
        conn,
        items,
        source=OFFICIAL_WEIGHT_SOURCE + " (12대 분류 시드)",
        valid_from="202001",
        # 2026-12 개편 전까지 유효하다. 개편분 적재 시 이 행은 닫히고 새 행이 열린다.
        valid_to=None,
    )


def load_from_db(conn: sqlite3.Connection, *, base_year: int) -> list[ItemWeight]:
    """적재된 가중치를 읽는다. 기준연도가 다른 행은 섞이지 않는다."""
    rows = conn.execute(
        "SELECT w.item_code, w.item_name, w.weight, w.base_year, m.division_code "
        "FROM item_weights w "
        "LEFT JOIN item_mapping m "
        "  ON m.item_code = w.item_code AND m.base_year = w.base_year "
        "WHERE w.base_year = ?",
        (base_year,),
    ).fetchall()
    return [
        ItemWeight(
            item_code=row["item_code"],
            item_name=row["item_name"],
            weight=row["weight"],
            division_code=row["division_code"] or row["item_code"],
            base_year=row["base_year"],
        )
        for row in rows
    ]


def division_totals(items: Sequence[ItemWeight]) -> dict[str, float]:
    """품목 가중치를 12대 분류로 합산한다."""
    totals: dict[str, float] = {}
    for item in items:
        totals[item.division_code] = totals.get(item.division_code, 0.0) + item.weight
    return totals


def close_validity(conn: sqlite3.Connection, *, base_year: int, valid_to: str) -> int:
    """기준개편 시 이전 기준연도 매핑의 유효기간을 닫는다."""
    cursor = conn.execute(
        "UPDATE item_mapping SET valid_to = ? WHERE base_year = ? AND valid_to IS NULL",
        (valid_to, base_year),
    )
    conn.commit()
    return cursor.rowcount


def loaded_at() -> str:
    return datetime.now(KST).isoformat()
